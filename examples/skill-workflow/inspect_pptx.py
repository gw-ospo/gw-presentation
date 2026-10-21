from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
import pptx.oxml.ns as ns
from lxml import etree

prs = Presentation('mnt/GW-brand/skill-workflow.pptx')
print(f"Slides: {len(prs.slides)}")
print(f"Width: {prs.slide_width.inches:.3f}\"  Height: {prs.slide_height.inches:.3f}\"")
print()

# Also check theme fonts
try:
    theme_xml = prs.slide_master.theme_color_map.element.getparent()
    print("Theme XML snippet (fonts):")
    print(etree.tostring(theme_xml, pretty_print=True).decode()[:2000])
except Exception as e:
    print(f"Theme error: {e}")

print()
print("=" * 70)

for slide_idx, slide in enumerate(prs.slides):
    print(f"\n=== SLIDE {slide_idx+1} ===")
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            for run in para.runs:
                text = run.text.strip()
                if not text:
                    continue
                rf = run.font
                pf = para.font

                name  = rf.name  or pf.name  or '(inherit)'
                size  = rf.size  or pf.size
                bold  = rf.bold
                italic = rf.italic
                try:
                    color = str(rf.color.rgb) if (rf.color and rf.color.type and rf.color.rgb) else None
                except Exception:
                    color = None
                size_pt = f"{size/12700:.1f}pt" if size else 'inherit'
                print(f"  \"{text[:55]}\"")
                print(f"    font={name}  size={size_pt}  bold={bold}  italic={italic}  color={color}")
